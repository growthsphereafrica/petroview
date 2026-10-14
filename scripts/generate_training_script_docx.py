"""
PetroView Forecourt OS — Master User Training Script & Facilitator's Guide Generator
Generates a complete, comprehensive, beautifully formatted Microsoft Word (.docx) document
and an accompanying Markdown (.md) document covering ALL aspects of the app for ALL users.
"""

import os
import sys
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

WORKSPACE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DOCS_DIR = os.path.join(WORKSPACE_DIR, 'docs')
os.makedirs(DOCS_DIR, exist_ok=True)

OUTPUT_DOCX = os.path.join(DOCS_DIR, 'PETROVIEW_MASTER_USER_TRAINING_SCRIPT.docx')
OUTPUT_ROOT_DOCX = os.path.join(WORKSPACE_DIR, 'PETROVIEW_MASTER_USER_TRAINING_SCRIPT.docx')
OUTPUT_MD = os.path.join(DOCS_DIR, 'PETROVIEW_MASTER_USER_TRAINING_SCRIPT.md')

# Color Palette Constants
COLOR_NAVY = RGBColor(30, 58, 138)        # #1E3A8A - Primary Header
COLOR_TEAL = RGBColor(14, 116, 144)      # #0E7490 - Secondary Header
COLOR_TEXT = RGBColor(30, 41, 59)        # #1E293B - Body Text
COLOR_MUTED = RGBColor(100, 116, 139)    # #64748B - Muted Subtitles
COLOR_EMERALD = RGBColor(5, 150, 105)    # #059669 - Hands-On Action
COLOR_AMBER = RGBColor(217, 119, 6)      # #D97706 - Warning / Pitfall
COLOR_INDIGO = RGBColor(79, 70, 229)     # #4F46E5 - Knowledge Check
COLOR_SKY = RGBColor(2, 132, 199)        # #0284C7 - Trainer Says
COLOR_WHITE = RGBColor(255, 255, 255)

HEX_NAVY = "1E3A8A"
HEX_TEAL = "0E7490"
HEX_SLATE_BG = "F8FAFC"
HEX_BORDER = "CBD5E1"

HEX_SAYS_BG = "F0F9FF"
HEX_SAYS_BORDER = "0284C7"

HEX_SHOWS_BG = "F8FAFC"
HEX_SHOWS_BORDER = "1E3A8A"

HEX_DO_BG = "ECFDF5"
HEX_DO_BORDER = "059669"

HEX_CHECK_BG = "EEF2FF"
HEX_CHECK_BORDER = "4F46E5"

HEX_WARN_BG = "FFFBEB"
HEX_WARN_BORDER = "D97706"

def set_cell_background(cell, hex_color):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=140, bottom=140, left=180, right=180):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, border_color=HEX_BORDER):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="single" w:sz="6" w:space="0" w:color="{border_color}"/>'
        f'  <w:bottom w:val="single" w:sz="8" w:space="0" w:color="{border_color}"/>'
        f'  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{border_color}"/>'
        f'  <w:insideV w:val="none"/>'
        f'  <w:left w:val="none"/>'
        f'  <w:right w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def set_callout_box_borders(table, border_color, border_sz=24):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="none"/>'
        f'  <w:bottom w:val="none"/>'
        f'  <w:left w:val="single" w:sz="{border_sz}" w:space="0" w:color="{border_color}"/>'
        f'  <w:right w:val="none"/>'
        f'  <w:insideH w:val="none"/>'
        f'  <w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

class TrainingDocBuilder:
    def __init__(self):
        self.doc = Document()
        self.md_lines = []
        self._setup_page()

    def _setup_page(self):
        for s in self.doc.sections:
            s.top_margin = Inches(0.8)
            s.bottom_margin = Inches(0.8)
            s.left_margin = Inches(0.8)
            s.right_margin = Inches(0.8)
            
            # Header
            hdr = s.header
            p_hdr = hdr.paragraphs[0]
            p_hdr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            r_hdr = p_hdr.add_run("PETROVIEW FORECOURT OS — MASTER TRAINING SCRIPT & FACILITATOR'S GUIDE")
            r_hdr.font.name = 'Segoe UI'
            r_hdr.font.size = Pt(8.5)
            r_hdr.font.color.rgb = COLOR_MUTED
            
            # Footer
            ftr = s.footer
            p_ftr = ftr.paragraphs[0]
            p_ftr.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r_ftr = p_ftr.add_run("CONFIDENTIAL & PROPRIETARY — FOR AUTHORIZED FACILITATORS & USERS · PETROVIEW OS V2.0")
            r_ftr.font.name = 'Segoe UI'
            r_ftr.font.size = Pt(8)
            r_ftr.font.color.rgb = COLOR_MUTED

    def add_title_block(self, title, subtitle, meta_items):
        # Docx Title
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(title)
        run.font.name = 'Segoe UI'
        run.font.size = Pt(22)
        run.font.bold = True
        run.font.color.rgb = COLOR_NAVY

        # Subtitle
        p_sub = self.doc.add_paragraph()
        p_sub.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p_sub.paragraph_format.space_after = Pt(12)
        r_sub = p_sub.add_run(subtitle)
        r_sub.font.name = 'Segoe UI'
        r_sub.font.size = Pt(12)
        r_sub.font.color.rgb = COLOR_TEAL

        # Meta Box Table
        tbl = self.doc.add_table(rows=len(meta_items), cols=2)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        for idx, (label, val) in enumerate(meta_items):
            cell_lbl = tbl.cell(idx, 0)
            cell_val = tbl.cell(idx, 1)
            set_cell_background(cell_lbl, HEX_SLATE_BG)
            set_cell_background(cell_val, HEX_SLATE_BG)
            set_cell_margins(cell_lbl, top=50, bottom=50, left=100, right=100)
            set_cell_margins(cell_val, top=50, bottom=50, left=100, right=100)
            
            p1 = cell_lbl.paragraphs[0]
            p1.paragraph_format.space_after = Pt(1)
            r1 = p1.add_run(label)
            r1.font.name = 'Segoe UI'
            r1.font.size = Pt(9)
            r1.font.bold = True
            r1.font.color.rgb = COLOR_NAVY
            
            p2 = cell_val.paragraphs[0]
            p2.paragraph_format.space_after = Pt(1)
            r2 = p2.add_run(val)
            r2.font.name = 'Segoe UI'
            r2.font.size = Pt(9)
            r2.font.color.rgb = COLOR_TEXT

        set_table_borders(tbl, HEX_BORDER)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(8)

        # Markdown Mirror
        self.md_lines.append(f"# {title}\n")
        self.md_lines.append(f"**{subtitle}**\n\n---\n")
        for label, val in meta_items:
            self.md_lines.append(f"- **{label}:** {val}")
        self.md_lines.append("\n---\n")

    def add_h1(self, text):
        h = self.doc.add_heading(level=1)
        h.paragraph_format.space_before = Pt(20)
        h.paragraph_format.space_after = Pt(5)
        run = h.add_run(text)
        run.font.name = 'Segoe UI'
        run.font.size = Pt(16)
        run.font.bold = True
        run.font.color.rgb = COLOR_NAVY

        self.md_lines.append(f"\n# {text}\n")

    def add_h2(self, text):
        h = self.doc.add_heading(level=2)
        h.paragraph_format.space_before = Pt(14)
        h.paragraph_format.space_after = Pt(4)
        run = h.add_run(text)
        run.font.name = 'Segoe UI'
        run.font.size = Pt(13)
        run.font.bold = True
        run.font.color.rgb = COLOR_TEAL

        self.md_lines.append(f"\n## {text}\n")

    def add_h3(self, text):
        h = self.doc.add_heading(level=3)
        h.paragraph_format.space_before = Pt(10)
        h.paragraph_format.space_after = Pt(2)
        run = h.add_run(text)
        run.font.name = 'Segoe UI'
        run.font.size = Pt(10.5)
        run.font.bold = True
        run.font.color.rgb = COLOR_TEXT

        self.md_lines.append(f"\n### {text}\n")

    def add_p(self, text, bold_prefix=None, space_after=4):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.15
        
        if bold_prefix:
            r_bold = p.add_run(bold_prefix + " ")
            r_bold.font.name = 'Segoe UI'
            r_bold.font.size = Pt(9.5)
            r_bold.font.bold = True
            r_bold.font.color.rgb = COLOR_NAVY
            
        r_text = p.add_run(text)
        r_text.font.name = 'Segoe UI'
        r_text.font.size = Pt(9.5)
        r_text.font.color.rgb = COLOR_TEXT

        if bold_prefix:
            self.md_lines.append(f"**{bold_prefix}** {text}\n")
        else:
            self.md_lines.append(f"{text}\n")

    def add_bullet(self, text, bold_prefix=None):
        p = self.doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(2.5)
        p.paragraph_format.line_spacing = 1.15
        
        if bold_prefix:
            r_bold = p.add_run(bold_prefix + " ")
            r_bold.font.name = 'Segoe UI'
            r_bold.font.size = Pt(9.5)
            r_bold.font.bold = True
            r_bold.font.color.rgb = COLOR_NAVY
            
        r_text = p.add_run(text)
        r_text.font.name = 'Segoe UI'
        r_text.font.size = Pt(9.5)
        r_text.font.color.rgb = COLOR_TEXT

        if bold_prefix:
            self.md_lines.append(f"- **{bold_prefix}** {text}")
        else:
            self.md_lines.append(f"- {text}")

    def add_table(self, headers, rows, col_widths=None):
        num_cols = len(headers)
        tbl = self.doc.add_table(rows=len(rows) + 1, cols=num_cols)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(tbl)

        # Header Row
        hdr_row = tbl.rows[0]
        for c_idx, h_text in enumerate(headers):
            cell = hdr_row.cells[c_idx]
            set_cell_background(cell, HEX_NAVY)
            set_cell_margins(cell, top=90, bottom=90, left=100, right=100)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(h_text)
            r.font.name = 'Segoe UI'
            r.font.size = Pt(9)
            r.font.bold = True
            r.font.color.rgb = COLOR_WHITE

        # Data Rows
        for r_idx, r_data in enumerate(rows):
            row = tbl.rows[r_idx + 1]
            bg = HEX_SLATE_BG if r_idx % 2 == 1 else "FFFFFF"
            for c_idx, val in enumerate(r_data):
                cell = row.cells[c_idx]
                set_cell_background(cell, bg)
                set_cell_margins(cell, top=70, bottom=70, left=100, right=100)
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(2)
                r = p.add_run(str(val))
                r.font.name = 'Segoe UI'
                r.font.size = Pt(8.5)
                r.font.color.rgb = COLOR_TEXT

        if col_widths and len(col_widths) == num_cols:
            for row in tbl.rows:
                for c_idx, w in enumerate(col_widths):
                    row.cells[c_idx].width = Inches(w)

        self.doc.add_paragraph().paragraph_format.space_after = Pt(5)

        # Markdown Table
        self.md_lines.append("\n| " + " | ".join(headers) + " |")
        self.md_lines.append("| " + " | ".join(["---"] * num_cols) + " |")
        for r_data in rows:
            self.md_lines.append("| " + " | ".join([str(v) for v in r_data]) + " |")
        self.md_lines.append("\n")

    def add_callout(self, badge_title, body_text_list, border_hex, bg_hex, badge_color_rgb):
        tbl = self.doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.cell(0, 0)
        set_cell_background(cell, bg_hex)
        set_callout_box_borders(tbl, border_hex, border_sz=24)
        set_cell_margins(cell, top=100, bottom=100, left=160, right=160)

        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(3)
        
        r_badge = p.add_run(f"[{badge_title.upper()}]\n")
        r_badge.font.name = 'Segoe UI'
        r_badge.font.size = Pt(9.5)
        r_badge.font.bold = True
        r_badge.font.color.rgb = badge_color_rgb

        for idx, line in enumerate(body_text_list):
            if idx > 0:
                p_item = cell.add_paragraph()
                p_item.paragraph_format.space_after = Pt(2.5)
                p_item.paragraph_format.line_spacing = 1.15
            else:
                p_item = p
                
            r = p_item.add_run(line)
            r.font.name = 'Segoe UI'
            r.font.size = Pt(9)
            r.font.color.rgb = COLOR_TEXT

        self.doc.add_paragraph().paragraph_format.space_after = Pt(5)

        # Markdown Mirror
        self.md_lines.append(f"\n> **{badge_title.upper()}**")
        for line in body_text_list:
            self.md_lines.append(f"> {line}")
        self.md_lines.append("\n")

    def add_trainer_says(self, speech_text):
        lines = [f'"{speech_text}"']
        self.add_callout("🗣️ FACILITATOR TALKING POINTS (SAY THIS)", lines, HEX_SAYS_BORDER, HEX_SAYS_BG, COLOR_SKY)

    def add_trainer_shows(self, steps_list):
        self.add_callout("🖥️ LIVE SYSTEM DEMONSTRATION (SHOW THIS ON PROJECTOR / SCREEN)", steps_list, HEX_SHOWS_BORDER, HEX_SHOWS_BG, COLOR_NAVY)

    def add_trainees_do(self, action_list):
        self.add_callout("✍️ HANDS-ON PARTICIPANT EXERCISE (TRAINEES DO THIS NOW)", action_list, HEX_DO_BORDER, HEX_DO_BG, COLOR_EMERALD)

    def add_knowledge_check(self, qa_list):
        formatted_lines = []
        for q, a in qa_list:
            formatted_lines.append(f"• Question: {q}")
            formatted_lines.append(f"  Expected Correct Answer: {a}\n")
        self.add_callout("❓ KNOWLEDGE CHECK & COMPREHENSION VERIFICATION", formatted_lines, HEX_CHECK_BORDER, HEX_CHECK_BG, COLOR_INDIGO)

    def add_pitfall_box(self, title, description, remedy):
        lines = [
            f"Common Field Mistake: {description}",
            f"Official Resolution & Recovery: {remedy}"
        ]
        self.add_callout(f"⚠️ FIELD PITFALL & RECOVERY: {title}", lines, HEX_WARN_BORDER, HEX_WARN_BG, COLOR_AMBER)

    def save(self):
        self.doc.save(OUTPUT_DOCX)
        self.doc.save(OUTPUT_ROOT_DOCX)
        with open(OUTPUT_MD, 'w', encoding='utf-8') as f:
            f.write('\n'.join(self.md_lines))
        print(f"Saved docx to: {OUTPUT_DOCX}")
        print(f"Saved root docx to: {OUTPUT_ROOT_DOCX}")
        print(f"Saved markdown to: {OUTPUT_MD}")

print("DocBuilder ready.")
